#!/usr/bin/env python3
"""Create and audit a shortcut-clean Harbor dataset split."""

from __future__ import annotations

import argparse
import csv
import errno
import gzip
import hashlib
import io
import json
import os
import shutil
import stat
import tempfile
import zipfile
from collections.abc import Callable
from contextlib import contextmanager
from pathlib import Path, PurePosixPath


TASK_PREFIX = "bbh-task__"
DEFAULT_MASK_PATH = Path(__file__).with_name("training_mask.txt")

DELETE_PATHS = {
    "015ab7f4-b069-4912-ac0b-a3f8bd9c869d": (
        "DE.csv",
        "README.md",
        "generate_ieeg_gradients.asv",
        "generate_ieeg_gradients.m",
        "gradients_and_networks.m",
        "gradients_multilinear.m",
        "title_fig.png",
    ),
    "79d5a5bc-0469-4a85-87d1-fe5d255b9823": ("ADBXD_GxE_manuscript analysis code.Rmd",),
    "08d9001f-42ba-49a2-bfb2-131a3898a97c": ("ADBXD_GxE_manuscript analysis code.Rmd",),
    "1d65c580-960c-4258-b987-1bc4d0469820": ("downloaded_only_Analysis_pamphilus.R",),
    "33ca12c3-5a9a-45f4-aded-305cc3c27f07": (
        "2-Model_Frac_eq.R",
        "2-RMA_comparison.txt",
        "downloaded_only_code_and_data.zip",
    ),
    "38a25e3c-1dd6-4b1f-b8ae-092134ec5ca9": (
        "downloaded_only_step1_preprocessing.m",
        "downloaded_only_step2b_markovchain.m",
    ),
    "3e09f25a-92ab-414f-9aee-a5c98f770c45": (
        "TCR signal strength modulates antigen-specific CD8+ T cell "
        "pathogenicity in non-obese diabetic mice.R",
    ),
    "624fca77-a490-48f2-87ff-78e6edea8217": ("downloaded_only_Fig04_plots.m",),
    "6fb44b38-5e45-4145-8c6c-9d6b283c8bb2": (
        "02_stats_mortalities_viralLoad_Fig3.R",
        "load_viral_moyen_TR463.csv",
        "pourcentages_mortalites_proxy_replicates.csv",
    ),
    "77f99441-0959-4031-84b9-de5d7c13c28e": ("downloaded_only_Cadmium_script.R",),
    "7eb68875-35ac-4723-beff-b996875dc612": (
        "TCR signal strength modulates antigen-specific CD8+ T cell "
        "pathogenicity in non-obese diabetic mice.R",
    ),
    "9ba1ec89-b67b-47a8-bd74-d4441bf1d718": (
        "6_SlingshotforTrajectory.R",
        "7_Sox10TimpointAnalysis",
    ),
    "98ca02a3-9cf7-4997-ae5a-9e8f939c01b0": ("2025.01.17.633577v4.full.pdf",),
    "a33a7c14-d960-4f5e-8d31-2fc41f9bd4f1": (
        "downloaded_only_Nanopore_targeted_sequencing_analysis.sh",
    ),
    "a4c16602-5520-41c4-8aa8-3e69ac62d806": ("downloaded_only_1_TCGA.R",),
    "a8350a95-b0d8-4e40-aeed-2bd297043086": ("mtb_hyp2_drug_sensitivity.py",),
    "b402da23-502b-4ebc-879f-30b3d87a3dc6": (
        "LFMM.R",
        "candidate_snps.txt",
    ),
    "ec4d9a27-3a39-4d77-88e2-95896f28251c": (
        "generate_ieeg_gradients.m",
        "gradients_and_networks.m",
    ),
    "dccb5380-8ddc-43e7-9610-cdfefa9bc0c0": (
        "TCGA_GBM_IDH_WT_RAP2A_by_subtype.csv",
        "downloaded_only_TCGA_GBM_IDH_WT_RAP2A_by_subtype.csv",
        "TCGA_GBM_RAP2A_expression_with_annotations.csv",
    ),
    "f52b991d-3d1f-4780-a453-25ddbcc8215d": (
        "time-domain-error-30sec-in-cm.pkl",
        "time-domain-error-xy.pkl",
    ),
    "6be3b69f-5284-4f0f-889e-4339cab31746": (
        "downloaded_only_PCA_PLINKPRUNED.eigenval",
        "downloaded_only_PCA_PLINKPRUNED.eigenvec",
    ),
    "1cf79c8c-fb8c-453c-8788-c8958ab6f152": (
        "S2A_Total_CHIP_x_VAFlt03.png",
        "S2A_Total_CHIP_x_VAFlt03.png.checksum",
    ),
    "2a8a40d4-05b0-4eec-8bd2-825f61fc9f5d": (
        "enrichMap_GO_Issy_ASXL1_fibro.tiff",
        "enrichMap_GO_Issy_ASXL1_fibro.tiff.checksum",
        "enrichMap_GO_Issy_ASXL1_fibro_smallfont.tiff",
        "enrichMap_GO_Issy_ASXL1_fibro_smallfont.tiff.checksum",
    ),
    "30b33e47-92d8-4372-a9f6-da32896493d0": (
        "CHIP VAF mean proportions.xlsx",
        "CHIP VAF mean proportions.xlsx.checksum",
    ),
    "33b801bb-9b47-4a0a-9314-05325c82fde7": (
        "FHAC_Issy-Paper1_2E_v2.ipynb",
        "FHAC_Issy-Paper1_2E_v2.ipynb.checksum",
        "enrichMap_GO_Issy_ASXL1_blood.tiff",
        "enrichMap_GO_Issy_ASXL1_blood.tiff.checksum",
        "enrichMap_GO_Issy_ASXL1_blood_smallfont.tiff",
        "enrichMap_GO_Issy_ASXL1_blood_smallfont.tiff.checksum",
    ),
    "7718a922-ce2c-4e59-900b-84fe06050ce6": (
        "S2B_CHIP_VAFlt03_Cons.png",
        "S2B_CHIP_VAFlt03_Cons.png.checksum",
    ),
}

DELETE_GLOBS = {
    "047a8b26-2bbe-415d-803e-903c4d6ed214": ("*.pdf",),
    "14289050-dcc7-43e7-a51e-a0f5b77303b8": ("*.R", "Figure_*"),
    "382f157e-0902-4c2d-9c08-ffe9abc4684e": (
        "*analysis*.ipynb",
        "coupled_simulations.m",
        "fig*.m",
    ),
    "6043f842-703a-42db-b184-d77bb5d8b5d7": ("6-*",),
    "66ff78c6-2792-4d63-90b0-91abdc5bf96d": ("*.R", "*.r", "*.sh", "*.pdf"),
    "720852ec-500d-407e-9135-502db964be39": (
        "EVO-25-0137R1_Supplementary_Material.docx",
        "FileS*_resultsselectionanalysis_*",
        "downloaded_only_PMC12687342.tar.gz",
        "qpaf190*",
    ),
    "752243b8-06dc-4040-aff7-c38725b9b9ac": (
        "*InfectionExperiment_PopulationTPCs*",
        "*InfectionFieldSurvey_Analysis*",
    ),
    "7c15efeb-83d5-4079-87a7-a971b2ab1a27": ("*.py",),
    "81cc0058-ec89-499a-b973-3402a286c0ce": (
        "*InfectionExperiment_PopulationTPCs*",
        "*InfectionFieldSurvey_Analysis*",
    ),
    "9d9e635a-cb6e-46c7-8da2-00b059d0a55f": ("*.R", "Figure_*"),
    "b92413f6-a7ea-4701-be98-6a17077c24cd": (
        "Fig*.ipynb",
        "*_p05.pickle",
        "analysis_pipeline.py",
        "cluster_stats.py",
    ),
    "bd55bc97-8bfb-434b-a732-c5fa58e93423": ("*.R", "*.r", "*.sh"),
    "be59f70b-e862-4173-993c-b4807c07b65f": ("*.R", "*.r", "*.sh"),
    "c8cd05f4-ac6e-40a9-9fd8-27d0922d06c1": ("*.pdf",),
    "dee048b6-83d7-47ec-9c6e-b69a912b159c": ("*.ipynb",),
    "e9ddbdbc-c74e-4371-b70a-6f4477cfec18": ("*.Rmd", "*.ipynb"),
    "f0807e09-0c08-4992-b7bc-b30a2cf30c46": ("*.R", "*.pdf"),
}

COPY_INPUTS = {
    "720852ec-500d-407e-9135-502db964be39": (
        (
            "52dad468-cc9e-4b61-9f6c-4e71faeaad64",
            "Data_Caecilian.zip",
            "Data_Caecilian.zip",
        ),
    ),
    "a8350a95-b0d8-4e40-aeed-2bd297043086": (
        (
            "4ee38ef1-ea48-4c6a-856a-ff7c7166413d",
            "GSE222412_rawCountMatrix.csv.gz",
            "GSE222412_rawCountMatrix.csv.gz",
        ),
    ),
    "4ef3fcd8-1c35-466f-9d93-49b92f4ea760": (
        (
            "2a8a40d4-05b0-4eec-8bd2-825f61fc9f5d",
            "Issy_ASXL1_fibro_featureCounts_GeneTable_final.txt",
            "Issy_ASXL1_fibro_featureCounts_GeneTable_final.txt",
        ),
        (
            "2a8a40d4-05b0-4eec-8bd2-825f61fc9f5d",
            "Issy_ASXL1_fibro_coldata_gender.xlsx",
            "Issy_ASXL1_fibro_coldata_gender.xlsx",
        ),
    ),
}

PATCH_REASONS = {
    "015ab7f4-b069-4912-ac0b-a3f8bd9c869d": (
        "The capsule contains the paper's analysis source, a precomputed embedding "
        "distance matrix, and repository documentation/figure that expose the "
        "requested workflow. Independent gradient, GD, MPC, principal-gradient, "
        "and spin-permutation inputs remain available."
    ),
    "a8350a95-b0d8-4e40-aeed-2bd297043086": (
        "The capsule contains only a complete expert solution. The matching raw "
        "GSE222412 count matrix is copied from another task before removal."
    ),
    "720852ec-500d-407e-9135-502db964be39": (
        "Published RELAX/PAML results and the paper expose the requested answer. "
        "A cleaned alignment/tree archive from the related caecilian task is "
        "copied in before those privileged materials are removed."
    ),
    "98ca02a3-9cf7-4997-ae5a-9e8f939c01b0": (
        "The paper and analysis/figure subtrees expose the longitudinal soapberry "
        "result. The cleaned archive retains the survey, cross, site, and image "
        "inputs."
    ),
    "1cf79c8c-fb8c-453c-8788-c8958ab6f152": (
        "The supplied figure visualizes the requested CHIP burden relationship. "
        "Per-sample variant calls, cohort metadata, and the CHIP gene list remain."
    ),
    "2a8a40d4-05b0-4eec-8bd2-825f61fc9f5d": (
        "The enrichment maps expose the requested fibroblast GO findings. Raw "
        "featureCounts, sample metadata, and gene annotations remain."
    ),
    "30b33e47-92d8-4372-a9f6-da32896493d0": (
        "The workbook contains the completed multinomial comparisons and exact "
        "p-values requested by the rubric. Per-sample variant calls remain."
    ),
    "33b801bb-9b47-4a0a-9314-05325c82fde7": (
        "The notebook states the conclusion and exact adjusted p-values and "
        "contains the complete expert analysis; enrichment maps expose the same "
        "result. Raw featureCounts and sample metadata remain."
    ),
    "7718a922-ce2c-4e59-900b-84fe06050ce6": (
        "The supplied figure shows the requested cohort-wise CHIP effect-type "
        "proportions. Per-sample variant calls and metadata remain."
    ),
}

DEFAULT_PATCH_REASON = (
    "Validated direct analysis or result shortcut; independent task inputs "
    "remain in the policy-visible capsule."
)

RPKM_DERIVED_COLUMNS = {
    "mock vs. HSV1_4h logFC",
    "mock vs. HSV1_4h logCPM",
    "mock vs. HSV1_4h PValue",
    "mock vs. HSV1_4h FDR",
    "mock vs. HSV1_8h logFC",
    "mock vs. HSV1_8h logCPM",
    "mock vs. HSV1_8h PValue",
    "mock vs. HSV1_8h FDR",
}

AD_BXD_DERIVED_COLUMNS = {
    "Predicted_AAO",
    "Sensorimotor_composite",
    "YM_6moNtgChow",
    "PS4_6moNtgChow",
    "CFM_6moNtgChow",
    "Weight_6moNtgChow",
    "YM_residuals",
    "PS4_residuals",
    "CFM_residuals",
    "Weight_residuals",
    "YM_resid_norm",
    "PS4_resid_norm",
    "CFM_resid_norm",
    "Composite_resid",
}


def read_mask(path: Path) -> set[str]:
    return {
        value
        for raw_line in path.read_text().splitlines()
        if (value := raw_line.partition("#")[0].strip())
    }


@contextmanager
def atomic_replacement(path: Path):
    descriptor, temporary_name = tempfile.mkstemp(
        dir=path.parent,
        prefix=f".{path.name}.",
        suffix=".tmp",
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        yield temporary
        os.chmod(temporary, stat.S_IMODE(path.stat().st_mode))
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def is_paml_path(name: str) -> bool:
    parts = PurePosixPath(name).parts
    return "Data4_Selection analysis" in parts and any(
        part in {"PAML", "._PAML"} for part in parts
    )


def patch_zip_bytes(contents: bytes) -> tuple[bytes, bool]:
    source_buffer = io.BytesIO(contents)
    target_buffer = io.BytesIO()
    changed = False
    with (
        zipfile.ZipFile(source_buffer, "r") as source,
        zipfile.ZipFile(target_buffer, "w") as target,
    ):
        for info in source.infolist():
            if is_paml_path(info.filename):
                changed = True
                continue
            data = source.read(info)
            if PurePosixPath(info.filename).name == "Data_Caecilian.zip":
                data, nested_changed = patch_zip_bytes(data)
                changed = changed or nested_changed
            target.writestr(info, data)
    return target_buffer.getvalue(), changed


def patch_zip(path: Path) -> None:
    contents, changed = patch_zip_bytes(path.read_bytes())
    if not changed:
        return
    with atomic_replacement(path) as temporary:
        temporary.write_bytes(contents)


def is_soapberry_analysis_path(name: str) -> bool:
    parts = PurePosixPath(name).parts
    return (
        "__MACOSX" in parts
        or "code" in parts
        or "Figures" in parts
        or PurePosixPath(name).name == ".DS_Store"
    )


def patch_soapberry_zip(path: Path) -> None:
    with zipfile.ZipFile(path, "r") as source:
        retained = [
            (info, source.read(info))
            for info in source.infolist()
            if not is_soapberry_analysis_path(info.filename)
        ]
    with atomic_replacement(path) as temporary:
        with zipfile.ZipFile(temporary, "w") as target:
            for info, contents in retained:
                target.writestr(info, contents)


def soapberry_zip_contains_analysis(path: Path) -> bool:
    with zipfile.ZipFile(path) as archive:
        return any(
            is_soapberry_analysis_path(info.filename) for info in archive.infolist()
        )


def zip_contains_paml(contents: bytes) -> bool:
    with zipfile.ZipFile(io.BytesIO(contents), "r") as archive:
        for info in archive.infolist():
            if is_paml_path(info.filename):
                return True
            if PurePosixPath(info.filename).name == "Data_Caecilian.zip":
                if zip_contains_paml(archive.read(info)):
                    return True
    return False


def read_delimited(path: Path, *, delimiter: str) -> list[list[str]]:
    opener = gzip.open if path.suffix == ".gz" else open
    with opener(path, "rt", newline="") as handle:
        return list(csv.reader(handle, delimiter=delimiter))


def write_delimited(path: Path, rows: list[list[str]], *, delimiter: str) -> None:
    with atomic_replacement(path) as temporary:
        if path.suffix == ".gz":
            with temporary.open("wb") as raw_handle:
                with gzip.GzipFile(
                    filename="",
                    mode="wb",
                    fileobj=raw_handle,
                    mtime=0,
                ) as gzip_handle:
                    with io.TextIOWrapper(gzip_handle, newline="") as text_handle:
                        writer = csv.writer(
                            text_handle,
                            delimiter=delimiter,
                            lineterminator="\n",
                        )
                        writer.writerows(rows)
        else:
            with temporary.open("w", newline="") as handle:
                writer = csv.writer(
                    handle,
                    delimiter=delimiter,
                    lineterminator="\n",
                )
                writer.writerows(rows)


def remove_columns(
    path: Path,
    *,
    delimiter: str,
    should_remove: Callable[[str], bool],
) -> None:
    rows = read_delimited(path, delimiter=delimiter)
    if not rows:
        raise ValueError(f"table is empty: {path}")
    remove_indexes = {
        index for index, name in enumerate(rows[0]) if should_remove(name)
    }
    if not remove_indexes:
        return
    filtered_rows = [
        [value for index, value in enumerate(row) if index not in remove_indexes]
        for row in rows
    ]
    write_delimited(path, filtered_rows, delimiter=delimiter)


def table_header(path: Path, *, delimiter: str) -> list[str]:
    rows = read_delimited(path, delimiter=delimiter)
    if not rows:
        raise ValueError(f"table is empty: {path}")
    return rows[0]


def patch_rpkm_tables(data_dir: Path) -> None:
    for name in (
        "downloaded_only_GSE243613_gene_rpkm_table.txt",
        "downloaded_only_GSE243613_gene_rpkm_table.txt.gz",
    ):
        path = data_dir / name
        if not path.is_file():
            raise FileNotFoundError(path)
        remove_columns(
            path,
            delimiter="\t",
            should_remove=RPKM_DERIVED_COLUMNS.__contains__,
        )


def patch_micos_workbook(data_dir: Path) -> None:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to patch the MICOS workbook") from exc

    path = data_dir / "media-2 (1).xlsx"
    workbook = openpyxl.load_workbook(path)
    changed = False
    for sheet_name in ("MICOS_LIVER_EUR", "MICOS_LIVER_AFR"):
        sheet = workbook[sheet_name]
        matches = [
            cell.column
            for row in sheet.iter_rows()
            for cell in row
            if cell.value == "Significance"
        ]
        if len(matches) > 1:
            raise ValueError(f"multiple Significance columns in {path}:{sheet_name}")
        if matches:
            sheet.delete_cols(matches[0])
            changed = True
    if changed:
        with atomic_replacement(path) as temporary:
            workbook.save(temporary)


def workbook_contains_significance(path: Path) -> bool:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to audit the MICOS workbook") from exc

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    return any(
        cell.value == "Significance"
        for sheet_name in ("MICOS_LIVER_EUR", "MICOS_LIVER_AFR")
        for row in workbook[sheet_name].iter_rows()
        for cell in row
    )


def is_ad_bxd_derived(name: str) -> bool:
    return (
        name in AD_BXD_DERIVED_COLUMNS
        or name.startswith("Predicted_AAO.")
        or name.startswith("Sensorimotor_composite.")
    )


def patch_ad_bxd_tables(data_dir: Path) -> None:
    for name in (
        "GxE_all_phenotypes_plus_resid_indiv_20240826_filteredforGxEstrains.csv",
        "GxE_all_phenotypes_plus_surv_resid_strainavgs_20240826_filteredforGxEstrains.csv",
    ):
        path = data_dir / name
        if not path.is_file():
            raise FileNotFoundError(path)
        remove_columns(
            path,
            delimiter=",",
            should_remove=is_ad_bxd_derived,
        )


def task_data_dir(split: Path, task_id: str) -> Path:
    return split / f"{TASK_PREFIX}{task_id}" / "environment" / "data"


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def patch_targets(split: Path) -> list[tuple[str, Path]]:
    targets: list[tuple[str, Path]] = []
    for task_id, relative_paths in DELETE_PATHS.items():
        data_dir = task_data_dir(split, task_id)
        targets.extend(
            (task_id, data_dir / relative_path)
            for relative_path in relative_paths
            if (data_dir / relative_path).is_file()
        )
    for task_id, patterns in DELETE_GLOBS.items():
        data_dir = task_data_dir(split, task_id)
        for pattern in patterns:
            targets.extend(
                (task_id, path) for path in data_dir.glob(pattern) if path.is_file()
            )

    fixed_targets = {
        "52dad468-cc9e-4b61-9f6c-4e71faeaad64": (
            "Data_Caecilian.zip",
            "downloaded_only_Data_Caecilian.zip",
        ),
        "bf14e7d3-2afb-495b-9e16-b5961623cc04": (
            "downloaded_only_GSE243613_gene_rpkm_table.txt",
            "downloaded_only_GSE243613_gene_rpkm_table.txt.gz",
        ),
        "eb5d2fbe-30ca-4cdd-8705-734248c66e92": ("media-2 (1).xlsx",),
        "79d5a5bc-0469-4a85-87d1-fe5d255b9823": (
            "GxE_all_phenotypes_plus_resid_indiv_20240826_filteredforGxEstrains.csv",
            "GxE_all_phenotypes_plus_surv_resid_strainavgs_20240826_filteredforGxEstrains.csv",
        ),
        "98ca02a3-9cf7-4997-ae5a-9e8f939c01b0": ("data_and_code.zip",),
    }
    for task_id, relative_paths in fixed_targets.items():
        data_dir = task_data_dir(split, task_id)
        targets.extend(
            (task_id, data_dir / relative_path)
            for relative_path in relative_paths
            if (data_dir / relative_path).is_file()
        )

    blast_id = "a33a7c14-d960-4f5e-8d31-2fc41f9bd4f1"
    blast_dir = task_data_dir(split, blast_id)
    targets.extend(
        (blast_id, path)
        for path in blast_dir.glob(
            "downloaded_modified_filtered_SRR319160*_blast_results.txt"
        )
    )
    return sorted(targets, key=lambda item: (item[0], str(item[1])))


def hardlink_or_copy(source: Path, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    try:
        os.link(source, destination)
    except OSError as exc:
        if exc.errno != errno.EXDEV:
            raise
        shutil.copy2(source, destination)


def quarantine_patch_targets(split: Path, destination: Path) -> None:
    if destination.exists():
        raise FileExistsError(f"privileged context destination exists: {destination}")
    destination.mkdir(parents=True)

    manifest_entries = []
    for task_id, source in patch_targets(split):
        relative_path = source.relative_to(task_data_dir(split, task_id))
        archived = destination / f"capsule_{task_id}" / relative_path
        hardlink_or_copy(source, archived)
        manifest_entries.append(
            {
                "archive_path": str(archived.relative_to(destination)),
                "capsule_uuid": task_id,
                "reason": PATCH_REASONS.get(task_id, DEFAULT_PATCH_REASON),
                "sha256": file_sha256(source),
                "size_bytes": source.stat().st_size,
                "source_path": str(relative_path),
            }
        )

    manifest = {
        "format_version": 1,
        "files": manifest_entries,
    }
    (destination / "manifest.json").write_text(
        json.dumps(manifest, indent=2, sort_keys=True) + "\n"
    )


def copy_replacement_inputs(split: Path) -> None:
    for target_id, inputs in COPY_INPUTS.items():
        for source_id, source_name, target_name in inputs:
            source = task_data_dir(split, source_id) / source_name
            target = task_data_dir(split, target_id) / target_name
            if not target.parent.is_dir():
                continue
            if not source.is_file():
                raise FileNotFoundError(source)
            target.unlink(missing_ok=True)
            hardlink_or_copy(source, target)


def apply_patches(split: Path) -> None:
    copy_replacement_inputs(split)
    for task_id, relative_paths in DELETE_PATHS.items():
        data_dir = task_data_dir(split, task_id)
        if not data_dir.is_dir():
            continue
        for relative_path in relative_paths:
            (data_dir / relative_path).unlink(missing_ok=True)
    for task_id, patterns in DELETE_GLOBS.items():
        data_dir = task_data_dir(split, task_id)
        for pattern in patterns:
            for path in data_dir.glob(pattern):
                if path.is_file():
                    path.unlink()

    for task_id in (
        "52dad468-cc9e-4b61-9f6c-4e71faeaad64",
        "720852ec-500d-407e-9135-502db964be39",
    ):
        caecilian = task_data_dir(split, task_id)
        if not caecilian.is_dir():
            continue
        for name in ("Data_Caecilian.zip", "downloaded_only_Data_Caecilian.zip"):
            path = caecilian / name
            if path.is_file() and zipfile.is_zipfile(path):
                patch_zip(path)

    rpkm = task_data_dir(split, "bf14e7d3-2afb-495b-9e16-b5961623cc04")
    if rpkm.is_dir():
        patch_rpkm_tables(rpkm)

    micos = task_data_dir(split, "eb5d2fbe-30ca-4cdd-8705-734248c66e92")
    if micos.is_dir():
        patch_micos_workbook(micos)

    ad_bxd = task_data_dir(split, "79d5a5bc-0469-4a85-87d1-fe5d255b9823")
    if ad_bxd.is_dir():
        patch_ad_bxd_tables(ad_bxd)

    soapberry = task_data_dir(split, "98ca02a3-9cf7-4997-ae5a-9e8f939c01b0")
    if soapberry.is_dir():
        patch_soapberry_zip(soapberry / "data_and_code.zip")

    blast = task_data_dir(split, "a33a7c14-d960-4f5e-8d31-2fc41f9bd4f1")
    if blast.is_dir():
        for path in blast.glob(
            "downloaded_modified_filtered_SRR319160*_blast_results.txt"
        ):
            path.unlink()


def collect_patch_findings(split: Path) -> list[str]:
    findings: list[str] = []
    for task_id, relative_paths in DELETE_PATHS.items():
        data_dir = task_data_dir(split, task_id)
        if not data_dir.is_dir():
            continue
        for relative_path in relative_paths:
            if (data_dir / relative_path).exists():
                findings.append(f"{task_id}: leaked file remains: {relative_path}")
    for task_id, patterns in DELETE_GLOBS.items():
        data_dir = task_data_dir(split, task_id)
        for pattern in patterns:
            for path in data_dir.glob(pattern):
                if path.is_file():
                    findings.append(
                        f"{task_id}: leaked glob match remains: {path.name}"
                    )

    caecilian_archives = {
        "52dad468-cc9e-4b61-9f6c-4e71faeaad64": (
            "Data_Caecilian.zip",
            "downloaded_only_Data_Caecilian.zip",
        ),
        "720852ec-500d-407e-9135-502db964be39": ("Data_Caecilian.zip",),
    }
    for task_id, names in caecilian_archives.items():
        caecilian = task_data_dir(split, task_id)
        if not caecilian.is_dir():
            continue
        for name in names:
            path = caecilian / name
            if not path.is_file():
                findings.append(f"{task_id}: missing {name}")
            elif zip_contains_paml(path.read_bytes()):
                findings.append(f"{task_id}: PAML subtree remains in {name}")

    for target_id, inputs in COPY_INPUTS.items():
        for _, _, target_name in inputs:
            target = task_data_dir(split, target_id) / target_name
            if task_data_dir(split, target_id).is_dir() and not target.is_file():
                findings.append(
                    f"{target_id}: replacement input is missing: {target_name}"
                )

    rpkm = task_data_dir(split, "bf14e7d3-2afb-495b-9e16-b5961623cc04")
    if rpkm.is_dir():
        for name in (
            "downloaded_only_GSE243613_gene_rpkm_table.txt",
            "downloaded_only_GSE243613_gene_rpkm_table.txt.gz",
        ):
            path = rpkm / name
            if not path.is_file():
                findings.append(f"bf14e7d3-2afb-495b-9e16-b5961623cc04: missing {name}")
                continue
            remaining = RPKM_DERIVED_COLUMNS.intersection(
                table_header(path, delimiter="\t")
            )
            if remaining:
                findings.append(
                    "bf14e7d3-2afb-495b-9e16-b5961623cc04: "
                    f"derived columns remain in {name}: {sorted(remaining)}"
                )

    micos = task_data_dir(split, "eb5d2fbe-30ca-4cdd-8705-734248c66e92")
    if micos.is_dir():
        path = micos / "media-2 (1).xlsx"
        if not path.is_file():
            findings.append("eb5d2fbe-30ca-4cdd-8705-734248c66e92: missing workbook")
        elif workbook_contains_significance(path):
            findings.append(
                "eb5d2fbe-30ca-4cdd-8705-734248c66e92: "
                "Significance column remains in workbook"
            )

    ad_bxd = task_data_dir(split, "79d5a5bc-0469-4a85-87d1-fe5d255b9823")
    if ad_bxd.is_dir():
        for name in (
            "GxE_all_phenotypes_plus_resid_indiv_20240826_filteredforGxEstrains.csv",
            "GxE_all_phenotypes_plus_surv_resid_strainavgs_20240826_filteredforGxEstrains.csv",
        ):
            path = ad_bxd / name
            if not path.is_file():
                findings.append(f"79d5a5bc-0469-4a85-87d1-fe5d255b9823: missing {name}")
                continue
            remaining = [
                value
                for value in table_header(path, delimiter=",")
                if is_ad_bxd_derived(value)
            ]
            if remaining:
                findings.append(
                    "79d5a5bc-0469-4a85-87d1-fe5d255b9823: "
                    f"derived columns remain in {name}: {remaining}"
                )

    soapberry = task_data_dir(split, "98ca02a3-9cf7-4997-ae5a-9e8f939c01b0")
    if soapberry.is_dir():
        path = soapberry / "data_and_code.zip"
        if not path.is_file():
            findings.append(
                "98ca02a3-9cf7-4997-ae5a-9e8f939c01b0: missing data archive"
            )
        elif soapberry_zip_contains_analysis(path):
            findings.append(
                "98ca02a3-9cf7-4997-ae5a-9e8f939c01b0: "
                "analysis or figure files remain in data_and_code.zip"
            )

    blast = task_data_dir(split, "a33a7c14-d960-4f5e-8d31-2fc41f9bd4f1")
    if blast.is_dir():
        leaked = sorted(
            path.name
            for path in blast.glob(
                "downloaded_modified_filtered_SRR319160*_blast_results.txt"
            )
        )
        if leaked:
            findings.append(
                "a33a7c14-d960-4f5e-8d31-2fc41f9bd4f1: "
                f"precomputed BLAST outputs remain: {leaked}"
            )
    return findings


def split_task_ids(split: Path) -> set[str]:
    return {
        path.name.removeprefix(TASK_PREFIX)
        for path in split.glob(f"{TASK_PREFIX}*")
        if path.is_dir()
    }


def prepare_split(
    source: Path,
    destination: Path,
    *,
    blocked_task_ids: set[str],
    privileged_context_dir: Path | None = None,
) -> None:
    if destination.exists():
        raise FileExistsError(f"destination already exists: {destination}")
    if source.resolve() == destination.resolve():
        raise ValueError("source and destination splits must differ")
    shutil.copytree(source, destination, copy_function=os.link)
    if privileged_context_dir is not None:
        quarantine_patch_targets(destination, privileged_context_dir)
    for task_id in blocked_task_ids:
        task_dir = destination / f"{TASK_PREFIX}{task_id}"
        if task_dir.exists():
            shutil.rmtree(task_dir)
    apply_patches(destination)

    findings = collect_patch_findings(destination)
    if findings:
        raise RuntimeError("shortcut patch audit failed:\n" + "\n".join(findings))
    remaining_blocked = split_task_ids(destination).intersection(blocked_task_ids)
    if remaining_blocked:
        raise RuntimeError(f"blocked tasks remain: {sorted(remaining_blocked)}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)

    audit = subparsers.add_parser("audit", help="Report shortcut leaks in a split.")
    audit.add_argument("split", type=Path)
    audit.add_argument(
        "--mask-file",
        type=Path,
        default=DEFAULT_MASK_PATH,
        help=f"task UUID mask to enforce (default: {DEFAULT_MASK_PATH})",
    )

    prepare = subparsers.add_parser(
        "prepare",
        help="Hard-link-clone, patch, and mask a source split.",
    )
    prepare.add_argument("source_split", type=Path)
    prepare.add_argument("destination_split", type=Path)
    prepare.add_argument(
        "--mask-file",
        type=Path,
        default=DEFAULT_MASK_PATH,
        help=f"task UUID mask to apply (default: {DEFAULT_MASK_PATH})",
    )
    prepare.add_argument(
        "--privileged-context-dir",
        type=Path,
        help=(
            "optional new directory that receives original shortcut files and a "
            "hash manifest before the clean split is patched"
        ),
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    blocked_task_ids = read_mask(args.mask_file.resolve())
    if args.command == "audit":
        split = args.split.resolve()
        findings = collect_patch_findings(split)
        blocked = split_task_ids(split).intersection(blocked_task_ids)
        for finding in findings:
            print(f"FAIL {finding}")
        print(
            f"audited {len(split_task_ids(split))} tasks: "
            f"{len(findings)} shortcut findings, {len(blocked)} blocked tasks present"
        )
        return 1 if findings or blocked else 0

    source = args.source_split.resolve()
    destination = args.destination_split.resolve()
    prepare_split(
        source,
        destination,
        blocked_task_ids=blocked_task_ids,
        privileged_context_dir=(
            args.privileged_context_dir.resolve()
            if args.privileged_context_dir is not None
            else None
        ),
    )
    print(
        f"prepared {len(split_task_ids(destination))} clean tasks at {destination}; "
        f"masked {len(split_task_ids(source).intersection(blocked_task_ids))} "
        "blocked tasks"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
